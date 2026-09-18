from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch, Circle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/workspace/scratch/6e022ae97052')
ASSETS = OUT / 'game_doc_assets'
ASSETS.mkdir(exist_ok=True)
DOCX = OUT / 'Project_Mayhem_Game_Design_Blueprint.docx'

NAVY = '142032'; TEAL = '1F8A89'; ORANGE = 'F28C45'; ICE = 'DDECEF'; LIGHT = 'F4F6F8'; DARK = '20252B'; RED = 'B94A48'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for tag, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+tag))
        if node is None: node = OxmlElement('w:'+tag); tcMar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr(); cant = OxmlElement('w:cantSplit'); cant.set(qn('w:val'),'true'); trPr.append(cant)

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    set_repeat_table_header(t.rows[0])
    prevent_row_split(t.rows[0])
    for ridx,row in enumerate(rows):
        new_row=t.add_row(); prevent_row_split(new_row); cells=new_row.cells
        for i,val in enumerate(row):
            cells[i].text=str(val); shade(cells[i], 'FFFFFF' if ridx%2==0 else LIGHT); margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(DARK)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

levels = [
 ('01','UNSCHEDULED TEST','Workshop / morning light','Movement, grab, Pack commands', [('START','Brief'),('TEST BAY','Movement'),('LIFT SHAFT','Climb'),('BROKEN LINE','Physics'),('ESCAPE','Chase'),('HUB RETURN','Reward')]),
 ('02','THE GROWING PROBLEM','Bioluminescent forest','Plant growth and flexible branches', [('RIDGE','Observe'),('ROOT TUNNEL','Puzzle'),('CANOPY','Skill'),('GREENHOUSE','Combine'),('OVERGROWTH','Chase'),('CORE','Choice')]),
 ('03','MAGNETIC PERSONALITY','Ancient mining station','Magnet glove and polarity', [('YARD','Learn'),('CRUSHER','Timing'),('SHAFT','Vertical'),('POLARITY LAB','Puzzle'),('TRAIN','Set-piece'),('VAULT','Mastery')]),
 ('04','DELIVERY ATTEMPT','Wind-carved canyon','Carry unstable energy cargo', [('DEPOT','Pickup'),('BRIDGE','Balance'),('MARKET','Comedy'),('CABLEWAY','Choice'),('STORM GAP','Chase'),('DELIVERY?','Twist')]),
 ('05','EVERYTHING IS FINE','Self-repairing factory','Chain reactions and misprint bots', [('INTAKE','Sneak'),('ASSEMBLY','Systems'),('QA LAB','Copies'),('FURNACE','Pressure'),('CONTROL','Rewire'),('COLLAPSE','Escape')]),
 ('06','MOTHER CLUCKZILLA','Forest-meets-factory arena','Three-stage environmental boss', [('NEST','Reveal'),('PURSUIT','Stage 1'),('ALARM','Stage 2'),('TREE GRID','Stage 3'),('HAY CART','Finish'),('AFTERMATH','Story')]),
 ('S1','ZERO MARGIN','Abandoned test track','Advanced movement speed trial', [('GATE','Start'),('WALL RUN','Chain'),('AIR GAP','Dash'),('DROP','Recover'),('FINAL LINE','Perfect'),('BOARD','Score')]),
 ('S2','THE WRONG SOLUTION','Prototype storage vault','Open-ended gadget puzzle', [('ENTRY','Inspect'),('OBJECT ROOM','Experiment'),('SPLIT','3 routes'),('MERGE','Consequences'),('LOCK','Final puzzle'),('ARCHIVE','Secret')]),
]

def level_diagram(level):
    num,title,env,mech,zones=level
    fig,ax=plt.subplots(figsize=(12,3.2),dpi=160); fig.patch.set_facecolor('#F4F6F8'); ax.set_facecolor('#F4F6F8'); ax.axis('off')
    xs=[0.5,2.35,4.2,6.05,7.9,9.75]; colors=['#1F8A89','#469B9A','#75AC9B','#D5A44C','#E27D42','#B94A48']
    ax.plot([0.8,10.05],[1.1,1.1],color='#73808C',lw=4,zorder=0)
    for i,((name,action),x,c) in enumerate(zip(zones,xs,colors)):
        h=0.8 + (i%3)*0.18
        box=FancyBboxPatch((x,0.7),1.35,h,boxstyle='round,pad=0.06,rounding_size=0.08',fc=c,ec='white',lw=2)
        ax.add_patch(box); ax.text(x+0.675,1.17,name,ha='center',va='center',fontsize=8,color='white',weight='bold'); ax.text(x+0.675,0.9,action,ha='center',va='center',fontsize=7,color='white')
        if i<5: ax.annotate('',xy=(x+1.7,1.12),xytext=(x+1.4,1.12),arrowprops=dict(arrowstyle='->',color='#56616B',lw=1.5))
    ax.text(.5,2.68,f'{num}  {title}',fontsize=15,weight='bold',color='#142032')
    ax.text(.5,2.35,f'{env}  |  {mech}',fontsize=8.5,color='#56616B')
    ax.text(.5,.25,'Main route',fontsize=7,color='#1F8A89'); ax.text(2.1,.25,'▲ skill branch',fontsize=7,color='#D08831'); ax.text(4.0,.25,'◆ puzzle branch',fontsize=7,color='#B94A48'); ax.text(8.6,.25,'Checkpoint before pressure peak',fontsize=7,color='#56616B')
    ax.set_xlim(0,11.7); ax.set_ylim(0,3)
    p=ASSETS/f'level_{num}.png'; plt.savefig(p,bbox_inches='tight',facecolor=fig.get_facecolor()); plt.close(fig); return p

diagram_paths=[level_diagram(l) for l in levels]

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal'].font.size=Pt(9); styles['Normal'].font.color.rgb=RGBColor.from_string(DARK)
for s,size,color in [('Title',30,NAVY),('Heading 1',20,NAVY),('Heading 2',13,TEAL),('Heading 3',10,ORANGE)]:
    styles[s].font.name='Aptos Display'; styles[s]._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); styles[s].font.size=Pt(size); styles[s].font.color.rgb=RGBColor.from_string(color); styles[s].font.bold=True

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(3); r=p.add_run('PROJECT MAYHEM'); r.bold=True; r.font.size=Pt(32); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Game Design Blueprint'); r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Characters  •  Level Layouts  •  Technical Architecture'); r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(DARK)
doc.add_paragraph('Working title and character names are placeholders pending trademark clearance. This blueprint defines the first production slice: a cinematic, funny 2.5D adventure built for web, iOS and Android.', style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_picture(str(diagram_paths[0]),width=Inches(6.95)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('1. Product vision',level=1)
add_table(doc,['Pillar','Production rule'],[
 ('Responsive hero','Precise movement; physics affects the world more than the controls.'),
 ('Cinematic 2.5D','3D characters and environments constrained to readable side-scrolling lanes.'),
 ('Systemic comedy','Gadgets, creatures and machinery create funny consequences during play.'),
 ('One memorable idea','Every level introduces one signature mechanic and ends by recombining it.'),
 ('Cross-platform','Shared gameplay project; platform-specific input, payments and account adapters.'),
 ('Ethical revenue','Free opening, full-game unlock, cosmetics and expansion worlds; no pay-to-win.'),
], [1.7,5.2])

doc.add_page_break()
doc.add_heading('2. Character plan',level=1)
add_table(doc,['Character','Visual silhouette','Personality and story function','Gameplay function'],[
 ('Bix - field engineer','Lean 5-head figure; asymmetrical field jacket; optical lens; mechanical glove; compact hard-shell backpack.','Capable but overconfident. Conceals mistakes and treats accidents as tests. The player sees the world through Bix.','Run, jump, wall-slide, dash, grab, throw and equip two gadgets.'),
 ('Pack - companion system','Credible modular backpack with camera lens, folding arms, status display and replaceable cartridges.','Anxious, precise and quietly competitive. Absorbs Chaos energy and becomes the final boss before reconciling.','Stores gadgets, scans objects, provides contextual actions and later supports co-op control.'),
 ('Luma - energy organism','Hand-sized translucent ribbon form with an illuminated inner structure and particle trail.','Excitable guide with poor direction sense. Communicates through movement, light and sound.','Detects secrets, activates ancient systems and provides optional hints.'),
 ('Vela - rival explorer','Athletic 5.5-head figure; clean technical suit; mechanical scarf/glider; precision equipment.','Disciplined rival whose perfect plans fail in imperfect worlds. Future playable character.','Appears in races; later offers speed-focused play with one gadget slot.'),
 ('Professor Ponk - mentor','Older engineer; layered workshop clothing; optical rig; leg brace; repaired coat.','Confident inventor responsible for the original accident. Delivers upgrades and unreliable history.','Tutorials, blueprints, workshop upgrades and side missions.'),
 ('Nib - operations lead','Small four-armed non-human; practical apron; inspection lenses; natural shell materials.','Serious production manager surrounded by chaos. Records every avoidable disaster.','Shop, cosmetics, crafting and mission administration.'),
 ('Baron Grumble - false villain','Tall 6.5-head armoured silhouette; weathered cloak; breathing unit; bronze and burgundy.','Looks dangerous but mainly wants quiet and theatrical respect. Eventually joins the heroes.','Recurring obstacle, story reversal and unlockable strength character.'),
], [1.2,1.8,2.7,1.8])

doc.add_heading('Character art rules',level=2)
for text in ['Use physically based fabric, brushed metal, worn polymer, dust and scratches.','Keep faces expressive but avoid oversized cartoon eyes or toy proportions.','Comedy comes from timing, posture, equipment behaviour and restrained facial acting.','Each principal character must remain identifiable as a black silhouette at mobile size.','Hero colours: deep teal, burnt orange, charcoal and warm-white illumination.','No red-cap/blue-overall plumber language, mushroom motifs, question blocks or Nintendo-like iconography.']:
    doc.add_paragraph(text,style='List Bullet')

doc.add_heading('3. First-world level layouts',level=1)
doc.add_paragraph('Each diagram shows dramatic flow, not final metric geometry. Every main level contains an accessible route, an optional skill branch, an optional gadget branch and a checkpoint before the highest-pressure section.')
for level,path in zip(levels,diagram_paths):
    num,title,env,mech,zones=level
    doc.add_heading(f'{num}. {title.title()}',level=2)
    doc.add_picture(str(path),width=Inches(6.9)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    if num=='01': notes='Teaches movement without text walls. Ends with a short machine-collapse chase and returns directly to the workshop.'
    elif num=='02': notes='Growth creates bridges, blocks routes and eventually becomes the pursuing threat. First major visual reveal of the living world.'
    elif num=='03': notes='Alternates slow polarity puzzles with fast industrial hazards. Magnet errors produce chain-reaction comedy.'
    elif num=='04': notes='The carried core changes weight and behaviour. Mid-level market scene provides NPC comedy and a low-pressure break.'
    elif num=='05': notes='All previous systems overlap. Player rewires a factory that continuously manufactures incorrect solutions.'
    elif num=='06': notes='Boss is solved through environment manipulation: pursuit, alarm reveal, awakened tree grid, then a non-violent hay-cart finish.'
    elif num=='S1': notes='Short mastery course with ghost replay, medals and leaderboard timing. No exclusive story progression.'
    else: notes='One open puzzle chamber with at least three valid solutions; records unusual solutions for challenge badges.'
    doc.add_paragraph(notes)

doc.add_heading('4. Whole-game architecture',level=1)
add_table(doc,['Layer','Responsibilities','Recommended implementation'],[
 ('Presentation','Website, account pages, store, news, support and embedded web build.','Next.js on Vercel'),
 ('Game runtime','Movement, physics, camera, animation, AI, level streaming, save client and UI.','Godot 4; GDScript/C#; web + native exports'),
 ('Content','Scenes, modular environment kits, level definitions, dialogue, audio and localisation.','Versioned Godot resources and addressable content manifests'),
 ('Platform adapters','Keyboard, touch, controller, Apple/Google sign-in, store purchases and notifications.','Interface-driven adapters per platform'),
 ('Backend','Authentication, profiles, cloud saves, entitlements, achievements and leaderboards.','Supabase Auth, PostgreSQL, Storage and Edge Functions'),
 ('Commerce','Receipt creation, verification, restoration and cross-device entitlements.','StoreKit, Google Play Billing and Stripe web checkout'),
 ('Operations','Crash reporting, analytics, remote configuration and limited experiments.','Privacy-conscious analytics; server-controlled config'),
], [1.25,3.15,2.65])

doc.add_heading('Runtime module map',level=2)
add_table(doc,['Module','Owns','Must not own'],[
 ('Game Flow','Boot, hub, level load, pause, completion and return.','Platform-specific billing logic.'),
 ('Player Controller','Movement state machine, abilities, health and interaction intent.','Direct save or network writes.'),
 ('Interaction System','Grab, throw, switches, gadget targets and context actions.','Character input mapping.'),
 ('Gadget System','Two-slot loadout, energy, combinations and controlled malfunctions.','Level-specific hard-coded exceptions.'),
 ('World Simulation','Physics objects, hazards, machines, creatures and checkpoints.','Account identity.'),
 ('Save Service','Local snapshot, versioning, offline queue, conflict resolution.','Raw passwords or payment credentials.'),
 ('Entitlement Service','Verified content access and restore state.','Trusting client purchase success.'),
 ('Telemetry','Events, performance and funnels using pseudonymous identifiers.','Unnecessary personal information.'),
], [1.45,3.2,2.4])

doc.add_heading('Data model',level=2)
add_table(doc,['Entity','Minimum fields'],[
 ('profiles','user_id, display_name, avatar_key, locale, created_at'),
 ('save_slots','user_id, slot, schema_version, progress_json, updated_at, device_id'),
 ('level_results','user_id, level_id, completed, best_time_ms, collectibles, route_badges'),
 ('inventory','user_id, item_id, source, acquired_at'),
 ('entitlements','user_id, product_id, provider, verified_at, status'),
 ('achievements','user_id, achievement_id, progress, unlocked_at'),
 ('leaderboard_entries','board_id, user_id, score, replay_hash, submitted_at'),
], [1.6,5.3])

doc.add_heading('5. Game loop and progression',level=1)
add_table(doc,['Loop','Player experience','Reward'],[
 ('Moment-to-moment','Move, observe, manipulate, improvise and recover from funny failures.','Mastery and spectacle'),
 ('Level','Learn one system, choose routes, combine systems, survive a set-piece.','Cogs, shard, blueprint progress'),
 ('World','Complete six levels, two secrets and a three-stage boss.','New gadget family and story chapter'),
 ('Long term','Complete challenges, improve times, collect cosmetics and expansions.','Badges, outfits, workshop growth'),
], [1.3,3.9,1.7])

doc.add_heading('6. Monetisation architecture',level=1)
add_table(doc,['Offer','Rule','Initial target'],[
 ('Free access','Workshop, tutorial and first three levels; no forced interstitial ads.','Acquisition and trust'),
 ('Full-world unlock','Permanent access, restorable across supported devices after verification.','AED 19.99-24.99 test'),
 ('Starter bundle','Full world, no ads and three cosmetic items.','AED 29.99 test'),
 ('Cosmetics','Visual only: outfits, Pack shells, trails, workshop decor.','AED 4.99-9.99'),
 ('Expansion world','Six levels, two secrets, gadgets, enemies and boss.','AED 14.99-19.99'),
 ('Rewarded ads','Optional revive, bonus coins or cosmetic trial; capped per session.','Free tier only'),
], [1.3,4.05,1.55])

doc.add_heading('7. Production plan',level=1)
add_table(doc,['Gate','Deliverable','Exit condition'],[
 ('G0 - Movement lab','Greybox room with keyboard, touch and controller input.','Movement remains satisfying for 15 minutes without content.'),
 ('G1 - Vertical slice','Level 01 with Bix, Pack, one enemy, one gadget and final chase.','Stable performance and positive blind-play feedback.'),
 ('G2 - Systems alpha','Save, login, level flow, three gadgets and reusable enemy framework.','No level-specific hacks in core modules.'),
 ('G3 - World alpha','Six levels, two secrets, hub and boss in greybox/final mix.','Complete start-to-boss playthrough.'),
 ('G4 - Soft launch','Purchases, analytics, cloud saves and limited-market mobile release.','Retention and conversion meet agreed targets.'),
 ('G5 - Release','Final art, audio, localisation, accessibility and store compliance.','Web, iOS and Android production builds approved.'),
], [1.35,3.3,2.25])

doc.add_heading('8. Immediate build backlog',level=1)
for i,text in enumerate([
 'Lock Bix, Pack and Luma silhouette sheets and material palette.',
 'Prototype Bix movement with coyote time, jump buffering, wall slide, dash and grab.',
 'Greybox Level 01 using the six-zone layout in this document.',
 'Implement one physics object, one machine, one enemy and one controlled malfunction.',
 'Create Pack lens/display reactions for success, danger and operator error.',
 'Run five blind playtests before producing Level 02.',
 'Only after movement approval, build authentication, cloud save and commerce adapters.',
]):
    doc.add_paragraph(f'{i+1}. {text}')

doc.add_heading('Prototype success criteria',level=2)
add_table(doc,['Metric','Target'],[
 ('First-time completion','At least 80% complete Level 01 without developer help.'),
 ('Control comprehension','Players understand movement and grabbing within 90 seconds.'),
 ('Memorable moment','At least 70% can describe the machine-collapse sequence afterward.'),
 ('Performance','60 FPS target; stable 30 FPS minimum on supported lower-tier phones.'),
 ('Replay intent','At least half voluntarily retry a branch, collectible or better time.'),
], [2.0,4.9])

footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=footer.add_run('PROJECT MAYHEM  |  GAME DESIGN BLUEPRINT  |  v0.1'); run.font.size=Pt(7); run.font.color.rgb=RGBColor(115,128,140)
doc.core_properties.title='Project Mayhem - Game Design Blueprint'; doc.core_properties.subject='Characters, level layouts and game architecture'; doc.core_properties.author='Rahul Sharma'
doc.save(DOCX)
print(DOCX)
